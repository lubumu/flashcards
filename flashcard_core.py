#!/usr/bin/env python3
"""Shared persistence and spaced repetition logic for Adeamus flashcards."""

import datetime
import os
import random
import re
import sqlite3
from collections import deque
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:
    raise RuntimeError("openpyxl is required to load the vocabulary Excel sheet.") from exc

DEFAULT_DB_PATH = Path("flashcards.db")
DEFAULT_EXCEL_PATH = Path("out/adeamus_flashcards.xlsx")


class DatabaseManager:
    """Manages the local SQLite database and Excel synchronization."""

    def __init__(self, db_path=DEFAULT_DB_PATH):
        self.db_path = str(db_path)
        self.init_db()

    def init_db(self):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS flashcards (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                lektion TEXT NOT NULL,
                latin TEXT NOT NULL,
                german TEXT NOT NULL,
                rating INTEGER DEFAULT 0,
                last_reviewed TEXT,
                review_count INTEGER DEFAULT 0,
                UNIQUE(lektion, latin)
            )
            """
        )
        conn.commit()
        conn.close()

    def _parse_headers(self, first_row):
        header = [str(cell).strip().lower() if cell is not None else "" for cell in first_row]
        try:
            return (
                header.index("lektion"),
                header.index("latein words"),
                header.index("german explainations"),
            )
        except ValueError:
            return (0, 1, 2)

    def _insert_or_update_flashcards(self, data_rows, lek_idx, lat_idx, ger_idx):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        added_count = 0

        for row in data_rows:
            if len(row) <= max(lek_idx, lat_idx, ger_idx):
                continue

            lektion = str(row[lek_idx]).strip() if row[lek_idx] is not None else ""
            latin = str(row[lat_idx]).strip() if row[lat_idx] is not None else ""
            german = str(row[ger_idx]).strip() if row[ger_idx] is not None else ""

            if not lektion or not latin or not german:
                continue

            cursor.execute(
                "SELECT id FROM flashcards WHERE lektion = ? AND latin = ?",
                (lektion, latin),
            )
            if cursor.fetchone():
                cursor.execute(
                    "UPDATE flashcards SET german = ? WHERE lektion = ? AND latin = ?",
                    (german, lektion, latin),
                )
            else:
                cursor.execute(
                    "INSERT INTO flashcards (lektion, latin, german, rating, review_count) VALUES (?, ?, ?, 0, 0)",
                    (lektion, latin, german),
                )
                added_count += 1

        conn.commit()
        conn.close()
        return added_count

    def sync_with_excel(self, excel_path=DEFAULT_EXCEL_PATH):
        excel_path = Path(excel_path)
        if not excel_path.exists():
            raise FileNotFoundError(f"Source vocabulary file not found at: {excel_path}")

        workbook = openpyxl.load_workbook(excel_path, read_only=True)
        sheet = workbook["flashcards"] if "flashcards" in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return 0

        lek_idx, lat_idx, ger_idx = self._parse_headers(rows[0])
        return self._insert_or_update_flashcards(rows[1:], lek_idx, lat_idx, ger_idx)

    def get_lessons(self):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute("SELECT DISTINCT lektion FROM flashcards")
        lessons = [row[0] for row in cursor.fetchall() if row[0]]
        conn.close()

        def natural_sort_key(value):
            return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value)]

        return sorted(lessons, key=natural_sort_key)

    def get_flashcards_by_lesson(self, lektion):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id, lektion, latin, german, rating, last_reviewed, review_count FROM flashcards WHERE lektion = ?",
            (lektion,),
        )
        columns = [col[0] for col in cursor.description]
        cards = [dict(zip(columns, row)) for row in cursor.fetchall()]
        conn.close()
        return cards

    def get_card(self, card_id):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id, lektion, latin, german, rating, last_reviewed, review_count FROM flashcards WHERE id = ?",
            (card_id,),
        )
        row = cursor.fetchone()
        columns = [col[0] for col in cursor.description] if cursor.description else []
        conn.close()
        return dict(zip(columns, row)) if row else None

    def update_rating(self, card_id, rating):
        now_str = datetime.datetime.now().isoformat()
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute(
            """
            UPDATE flashcards
            SET rating = ?,
                last_reviewed = ?,
                review_count = review_count + 1
            WHERE id = ?
            """,
            (rating, now_str, card_id),
        )
        conn.commit()
        conn.close()

    def get_lesson_stats(self, lektion=None):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        if lektion:
            cursor.execute(
                "SELECT COUNT(*), SUM(CASE WHEN rating >= 3 THEN 1 ELSE 0 END), AVG(CASE WHEN rating > 0 THEN rating ELSE NULL END) FROM flashcards WHERE lektion = ?",
                (lektion,),
            )
        else:
            cursor.execute(
                "SELECT COUNT(*), SUM(CASE WHEN rating >= 3 THEN 1 ELSE 0 END), AVG(CASE WHEN rating > 0 THEN rating ELSE NULL END) FROM flashcards"
            )
        total, learned, avg_rating = cursor.fetchone()
        conn.close()

        total = total or 0
        learned = learned or 0
        avg_rating = round(avg_rating, 2) if avg_rating is not None else 0.0
        progress_pct = round((learned / total) * 100, 1) if total > 0 else 0.0
        return {
            "total": total,
            "learned": learned,
            "avg_rating": avg_rating,
            "progress_pct": progress_pct,
        }

    def get_difficult_words(self, limit=5, lektion=None):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        if lektion:
            cursor.execute(
                """
                SELECT lektion, latin, german, rating, review_count
                FROM flashcards
                WHERE lektion = ? AND rating > 0 AND rating < 3
                ORDER BY rating ASC, review_count DESC
                LIMIT ?
                """,
                (lektion, limit),
            )
        else:
            cursor.execute(
                """
                SELECT lektion, latin, german, rating, review_count
                FROM flashcards
                WHERE rating > 0 AND rating < 3
                ORDER BY rating ASC, review_count DESC
                LIMIT ?
                """,
                (limit,),
            )
        rows = cursor.fetchall()
        conn.close()
        return rows

    def get_recent_activity(self, limit=10, lektion=None):
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        if lektion:
            cursor.execute(
                """
                SELECT lektion, latin, german, rating, last_reviewed, review_count
                FROM flashcards
                WHERE lektion = ? AND last_reviewed IS NOT NULL
                ORDER BY last_reviewed DESC
                LIMIT ?
                """,
                (lektion, limit),
            )
        else:
            cursor.execute(
                """
                SELECT lektion, latin, german, rating, last_reviewed, review_count
                FROM flashcards
                WHERE last_reviewed IS NOT NULL
                ORDER BY last_reviewed DESC
                LIMIT ?
                """,
                (limit,),
            )
        rows = cursor.fetchall()
        conn.close()
        return rows


class SpacedRepetitionEngine:
    """Adaptive card selector with weighted repetition and recency blocking."""

    def __init__(self, db_manager, lektion, recent_card_ids=None):
        self.db_manager = db_manager
        self.lektion = lektion
        self.reload_cards()
        recency_size = max(0, min(4, len(self.cards) // 2))
        self.recency_queue = deque(recent_card_ids or [], maxlen=recency_size)

    def reload_cards(self):
        self.cards = self.db_manager.get_flashcards_by_lesson(self.lektion)

    def _weight_for_rating(self, rating):
        if rating in (0, 1):
            return 10.0
        if rating == 2:
            return 5.0
        if rating == 3:
            return 2.0
        if rating == 4:
            return 0.5
        if rating == 5:
            return 0.1
        return 1.0

    def get_next_card(self):
        if not self.cards:
            return None

        self.reload_cards()
        card_candidates = []
        weights = []
        recent_ids = set(self.recency_queue)

        for card in self.cards:
            weight = 0.0 if card["id"] in recent_ids else self._weight_for_rating(card["rating"])
            card_candidates.append(card)
            weights.append(weight)

        if sum(weights) == 0:
            weights = [1.0 for _ in card_candidates]

        selected_card = random.choices(card_candidates, weights=weights, k=1)[0]
        self.recency_queue.append(selected_card["id"])
        return selected_card

    def recent_card_ids(self):
        return list(self.recency_queue)


def initialize_database(base_dir=None, db_path=None, excel_path=None):
    """Ensure the database exists and is synchronized with the workbook."""
    root = Path(base_dir or Path(__file__).resolve().parent)
    resolved_db = Path(db_path) if db_path else root / DEFAULT_DB_PATH
    resolved_excel = Path(excel_path) if excel_path else root / DEFAULT_EXCEL_PATH

    db_manager = DatabaseManager(db_path=resolved_db)
    if resolved_excel.exists():
        db_manager.sync_with_excel(resolved_excel)
    return db_manager
